"""Aplica a formatação acadêmica essencial ao DOCX oficial do TCC."""

from __future__ import annotations

import os
import re
import shutil
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt

ROOT = Path(__file__).resolve().parents[1]
DOCX = ROOT / "tcc" / "TCC_COMPLETO_Joao_Victor.docx"
BACKUP = ROOT / "tcc" / "TCC_COMPLETO_Joao_Victor_ANTES_FORMATACAO_20261015.docx"

PRIMARY_HEADING = re.compile(r"^[1-5] [A-ZÁÉÍÓÚÂÊÔÃÕÇ]")
NUMBERED_HEADING = re.compile(r"^[1-5](?:\.\d+)+ ")
CAPTION = re.compile(r"^(Figura|Quadro) \d+ ")


def set_run_font(run, size: int = 12, *, bold: bool | None = None) -> None:
    run.font.name = "Arial"
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{attribute}"), "Arial")


def set_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.paragraph_format.first_line_indent = Cm(0)
    run = paragraph.add_run()
    set_run_font(run, 10)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instruction, separate, text, end))


def has_page_field(section) -> bool:
    xml = section.header._element.xml
    return " PAGE " in xml or ">PAGE<" in xml


def main() -> None:
    if not BACKUP.exists():
        shutil.copy2(DOCX, BACKUP)

    document = Document(DOCX)

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(12)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:cs"), "Arial")
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.first_line_indent = Cm(1.25)

    for section in document.sections:
        section.page_width = Mm(210)
        section.page_height = Mm(297)
        section.top_margin = Cm(3)
        section.left_margin = Cm(2.5)
        section.bottom_margin = Cm(2)
        section.right_margin = Cm(2.5)
        section.header_distance = Cm(1.5)
        section.footer_distance = Cm(1.0)
        section.start_type = WD_SECTION.CONTINUOUS
        if not has_page_field(section):
            header = section.header
            paragraph = header.paragraphs[0]
            if paragraph.text:
                paragraph = header.add_paragraph()
            set_page_number(paragraph)

    references = False
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        paragraph.paragraph_format.widow_control = True
        paragraph.paragraph_format.space_after = Pt(0)
        for run in paragraph.runs:
            set_run_font(run, 12)

        if not text:
            paragraph.paragraph_format.first_line_indent = Cm(0)
            continue

        if text == "REFERÊNCIAS":
            references = True
            for run in paragraph.runs:
                set_run_font(run, 12, bold=True)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph.paragraph_format.space_before = Pt(12)
            paragraph.paragraph_format.space_after = Pt(12)
            paragraph.paragraph_format.page_break_before = True
            paragraph.paragraph_format.keep_with_next = True
            continue

        if references:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            paragraph.paragraph_format.space_after = Pt(12)
            continue

        if PRIMARY_HEADING.match(text) or NUMBERED_HEADING.match(text):
            for run in paragraph.runs:
                set_run_font(run, 12, bold=True)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph.paragraph_format.space_before = Pt(12)
            paragraph.paragraph_format.space_after = Pt(6)
            paragraph.paragraph_format.keep_with_next = True
            if PRIMARY_HEADING.match(text) and not text.startswith("1 "):
                paragraph.paragraph_format.page_break_before = True
            continue

        if CAPTION.match(text):
            for run in paragraph.runs:
                set_run_font(run, 10)
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER
                if text.startswith("Figura")
                else WD_ALIGN_PARAGRAPH.LEFT
            )
            paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            paragraph.paragraph_format.space_before = Pt(6)
            paragraph.paragraph_format.space_after = Pt(4)
            paragraph.paragraph_format.keep_with_next = True
            continue

        if text.startswith("Fonte:"):
            for run in paragraph.runs:
                set_run_font(run, 10)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            paragraph.paragraph_format.space_before = Pt(4)
            paragraph.paragraph_format.space_after = Pt(6)
            continue

        if paragraph._p.xpath(".//w:drawing"):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            paragraph.paragraph_format.space_after = Pt(0)
            continue

        if paragraph.style.name.startswith("List"):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph.paragraph_format.left_indent = Cm(1.25)
            paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
            continue

        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        paragraph.paragraph_format.first_line_indent = Cm(1.25)
        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE

    for table in document.tables:
        table.autofit = True
        for row_index, row in enumerate(table.rows):
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for paragraph in cell.paragraphs:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    paragraph.paragraph_format.first_line_indent = Cm(0)
                    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
                    paragraph.paragraph_format.space_after = Pt(0)
                    for run in paragraph.runs:
                        set_run_font(run, 10, bold=True if row_index == 0 else None)

    tmp = DOCX.with_suffix(".formatando.docx")
    document.save(tmp)
    os.replace(tmp, DOCX)


if __name__ == "__main__":
    main()
