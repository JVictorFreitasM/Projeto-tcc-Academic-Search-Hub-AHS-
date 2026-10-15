# ruff: noqa: E501, RUF001
"""Aplica correções bibliográficas e factuais ao DOCX oficial do TCC.

O script cria uma cópia de segurança antes da primeira alteração e grava o
resultado de forma atômica. Ele não depende de chaves de API.
"""

from __future__ import annotations

import os
import shutil
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.shared import Mm

ROOT = Path(__file__).resolve().parents[1]
DOCX = ROOT / "tcc" / "TCC_COMPLETO_Joao_Victor.docx"
BACKUP = ROOT / "tcc" / "TCC_COMPLETO_Joao_Victor_ANTES_REFERENCIAS_20261015.docx"


def find_paragraph(document: Document, prefix: str):
    matches = [p for p in document.paragraphs if p.text.startswith(prefix)]
    if len(matches) != 1:
        raise RuntimeError(
            f"Esperava um parágrafo iniciado por {prefix!r}; encontrei {len(matches)}"
        )
    return matches[0]


def replace_paragraph(document: Document, prefix: str, text: str) -> None:
    paragraph = find_paragraph(document, prefix)
    paragraph.text = text


def remove_paragraph(paragraph) -> None:
    element = paragraph._element
    element.getparent().remove(element)
    paragraph._p = paragraph._element = None


def insert_after(paragraph, text: str):
    new_paragraph = paragraph._parent.add_paragraph()
    new_paragraph.style = paragraph.style
    new_paragraph.text = text
    paragraph._p.addnext(new_paragraph._p)
    return new_paragraph


def append_source_rows(document: Document) -> None:
    table = document.tables[0]
    existing = {row.cells[0].text for row in table.rows}
    rows = (
        (
            "UFMG",
            "OAI-PMH, candidata ao índice local",
            "Identify válido; 100 registros analisados e dry-run de 50 registros concluído",
        ),
        (
            "UFSCar",
            "OAI-PMH, candidata ao índice local",
            "Identify válido; 100 registros analisados e dry-run de 50 registros concluído",
        ),
        (
            "Unesp",
            "OAI-PMH, candidata ao índice local",
            "Identify válido; 100 registros analisados e dry-run de 50 registros concluído",
        ),
        (
            "UFRGS",
            "OAI-PMH, candidata ao índice local",
            "Identify válido; 100 registros analisados e dry-run de 50 registros concluído; poucos resumos na amostra",
        ),
    )
    for values in rows:
        if values[0] in existing:
            continue
        cells = table.add_row().cells
        for cell, value in zip(cells, values, strict=True):
            cell.text = value


def main() -> None:
    if not BACKUP.exists():
        shutil.copy2(DOCX, BACKUP)

    document = Document(DOCX)

    for section in document.sections:
        section.page_width = Mm(210)
        section.page_height = Mm(297)
        section.start_type = WD_SECTION.CONTINUOUS

    replace_paragraph(
        document,
        "Servidor. Segundo Van Rossum",
        "Servidor. Van Rossum e De Boer (1991) apresentaram Python como uma linguagem interpretada e extensível, usada para desenvolver aplicações de maneira interativa. Menezes (2010) apresenta classes, objetos e exemplos de programação orientada a objetos em Python. Conforme a documentação oficial de Ramírez (s. d.), o FastAPI fornece validação de dados, serialização em JSON, documentação interativa no padrão OpenAPI e suporte a operações assíncronas.",
    )
    replace_paragraph(
        document,
        "Banco de dados. Stonebraker",
        "Banco de dados. Stonebraker e Kemnitz (1991) descrevem o POSTGRES, origem do PostgreSQL, como um sistema relacional estendido com suporte a tipos de dados complexos e regras. Jung et al. (2015) comparam o desempenho de entrada e saída de dados do MongoDB e do PostgreSQL em ambiente de big data, e Krocz, Kizun e Skublewska-Paszkowska (2020) analisam o desempenho dos bancos relacionais MySQL, PostgreSQL, MariaDB e H2.",
    )
    replace_paragraph(
        document,
        "Interface. O HTML5",
        "Interface. O HTML5 define os elementos semânticos que estruturam o conteúdo (W3C, 2014), e as consultas de mídia permitem adaptar a apresentação às características da tela e do dispositivo (W3C, 2012). Segundo Flanagan (2013), o JavaScript é a linguagem padrão para a lógica interativa no navegador, com manipulação do DOM e requisições assíncronas sem recarregar a página. Teruel (2014) apresenta de forma prática o uso do HTML5 na construção de interfaces web.",
    )
    replace_paragraph(
        document,
        "Conforme Bader e Hammouri",
        "Conforme Bader e Hammouri (2016), o design responsivo reúne técnicas que permitem à interface se adaptar a diferentes tamanhos de tela, resoluções e dispositivos, mantendo a experiência de uso consistente.",
    )
    replace_paragraph(
        document,
        "Quanto à acessibilidade, as Diretrizes",
        "Quanto à acessibilidade, as Diretrizes de Acessibilidade para Conteúdo Web (WCAG) 2.1 definem critérios para que o conteúdo possa ser percebido, operado e compreendido por pessoas com deficiências visuais, auditivas, motoras e cognitivas (W3C, 2018). Nagaraju, Chawla e Rana (2019) discutem os desafios práticos de avaliar a acessibilidade de sites segundo as WCAG 2.0.",
    )
    replace_paragraph(
        document,
        "Quanto à acessibilidade, a interface foi desenvolvida",
        find_paragraph(
            document, "Quanto à acessibilidade, a interface foi desenvolvida"
        ).text.replace("(W3C, 2018b)", "(W3C, 2018)"),
    )
    replace_paragraph(
        document,
        "A usabilidade foi avaliada por inspeção",
        "A usabilidade foi avaliada por inspeção, pelo método da avaliação heurística apresentado por Prates e Barbosa (2003) e discutido por Cybis, Betiol e Faust (2010), aplicando-se a cada tela as dez heurísticas de usabilidade de Nielsen (1994 apud PRATES; BARBOSA, 2003). A inspeção foi realizada sem a participação de usuários. Foram registrados 13 problemas: dois de gravidade 3, três de gravidade 2 e oito de gravidade 1. Os dois mais graves foram a perda da sessão ao voltar à busca depois de entrar, que devolvia o usuário à tela de entrada, e a ausência do contato do responsável na política de privacidade; entre os de gravidade 2 esteve a falta de aviso de que não há recuperação de senha. Os problemas de gravidade 3 e 2 foram corrigidos ou atenuados, e os de gravidade 1 foram aceitos com justificativa. A recuperação de senha não foi implementada, porque exigiria o envio de e-mails.",
    )

    source_paragraph = find_paragraph(document, "A verificação de cada fonte com respostas reais")
    extra = insert_after(
        source_paragraph,
        "Como verificação adicional da capacidade de expansão, foram testados em 15 de outubro de 2026 os repositórios da Universidade Federal de Minas Gerais (UFMG), da Universidade Federal de São Carlos (UFSCar), da Universidade Estadual Paulista (Unesp) e da Universidade Federal do Rio Grande do Sul (UFRGS). Os quatro responderam ao Identify do OAI-PMH, devolveram uma página de 100 registros em Dublin Core e concluíram uma simulação de colheita de 50 registros sem gravação no banco. UFMG, UFSCar e Unesp apresentaram boa cobertura de resumos na amostra; na UFRGS, apenas um dos 100 registros continha resumo. As quatro fontes foram acrescentadas à configuração como candidatas desabilitadas, para que a validação não ampliasse silenciosamente o escopo da busca nem iniciasse uma colheita de grande volume.",
    )
    extra.style = source_paragraph.style
    append_source_rows(document)

    replace_paragraph(
        document,
        "Das seis fontes previstas",
        "Das seis fontes previstas, três foram integradas e validadas com respostas reais: CORE, OpenAlex e BDTD, esta pela API REST, porque o seu endereço OAI-PMH não respondeu. SciELO, UFPI e UESPI não puderam ser integrados por razões externas à aplicação: no SciELO, nenhum endereço OAI-PMH respondeu; na UFPI e na UESPI, o índice OAI-PMH estava vazio e a cadeia de certificados estava incompleta. A colheita por OAI-PMH e a busca no índice local foram comprovadas com o repositório da UFPE, incluído para essa validação. UFMG, UFSCar, Unesp e UFRGS também tiveram o protocolo e uma amostra de registros validados, mas permaneceram desabilitadas por estarem fora do escopo originalmente definido. Assim, a busca nos repositórios da UESPI e da UFPI, parte central da motivação do trabalho, não foi alcançada nesta versão e depende da correção desses repositórios pelos seus administradores.",
    )
    replace_paragraph(
        document,
        "Novos repositórios:",
        "Novos repositórios: após avaliação de volume, qualidade dos metadados e aderência ao escopo, habilitar UFMG, UFSCar, Unesp e UFRGS e avaliar outros repositórios de universidades estaduais e federais.",
    )

    references = {
        "ALMEIDA,": "ALMEIDA, L. S.; DIAS, T. M. R. Análise dos mecanismos de busca de repositórios institucionais de acesso aberto. RDBCI: Revista Digital de Biblioteconomia e Ciência da Informação, Campinas, v. 21, e023008, 2023. DOI 10.20396/rdbci.v21i00.8672274. Disponível em: https://periodicos.sbu.unicamp.br/ojs/index.php/rdbci/article/view/8672274. Acesso em: 15 out. 2026.",
        "BRASIL.": "BRASIL. Lei nº 13.709, de 14 de agosto de 2018. Lei Geral de Proteção de Dados Pessoais (LGPD). Diário Oficial da União: seção 1, Brasília, DF, 15 ago. 2018. Disponível em: https://www.planalto.gov.br/ccivil_03/_ato2015-2018/2018/lei/l13709.htm. Acesso em: 15 out. 2026.",
        "BRAY,": "BRAY, T. (ed.). The JavaScript Object Notation (JSON) data interchange format. RFC 8259. [S. l.]: Internet Engineering Task Force, dez. 2017. Disponível em: https://www.rfc-editor.org/rfc/rfc8259. Acesso em: 15 out. 2026.",
        "CYBIS,": "CYBIS, W.; BETIOL, A. H.; FAUST, R. Ergonomia e usabilidade: conhecimentos, métodos e aplicações. 2. ed. São Paulo: Novatec, 2010.",
        "FLANAGAN,": "FLANAGAN, D. JavaScript: o guia definitivo. Tradução João Eduardo Nóbrega Tortello. 6. ed. Porto Alegre: Bookman, 2013.",
        "IBICT.": "IBICT. Manifesto brasileiro de apoio ao acesso livre à informação científica. Brasília, DF: IBICT, 2005. Disponível em: https://www.icict.fiocruz.br/sites/www.icict.fiocruz.br/files/Manifesto%20sobre%20o%20acesso%20livre%20a%20informacao%20cientifica.pdf. Acesso em: 15 out. 2026.",
        "JUNG,": "JUNG, M.-G.; YOUN, S.-A.; BAE, J.; CHOI, Y.-L. A study on data input and output performance comparison of MongoDB and PostgreSQL in the big data environment. In: INTERNATIONAL CONFERENCE ON DATABASE THEORY AND APPLICATION (DTA), 8., 2015. Proceedings [...]. [S. l.]: IEEE, 2015. p. 14-17. DOI 10.1109/DTA.2015.14.",
        "LYNCH,": "LYNCH, C. A. Institutional repositories: essential infrastructure for scholarship in the digital age. ARL Bimonthly Report, Washington, DC, n. 226, p. 1-7, fev. 2003. Disponível em: https://www.ischool.berkeley.edu/research/publications/2003/institutional-repositories-essential-infrastructure-scholarship-digital. Acesso em: 15 out. 2026.",
        "MENEZES,": "MENEZES, N. N. C. Introdução à programação com Python: algoritmos e lógica de programação para iniciantes. 1. ed. São Paulo: Novatec, 2010.",
        "NAGARAJU,": "NAGARAJU, M.; CHAWLA, P.; RANA, A. A practitioner's approach to assess the WCAG 2.0 website accessibility challenges. In: AMITY INTERNATIONAL CONFERENCE ON ARTIFICIAL INTELLIGENCE (AICAI), 2019. Proceedings [...]. [S. l.]: IEEE, 2019. p. 958-966. DOI 10.1109/AICAI.2019.8701320.",
        "NUNES,": "NUNES, P.; MEDEIROS, I.; FONSECA, J.; NEVES, N. F.; CORREIA, M.; VIEIRA, M. Benchmarking static analysis tools for web security. IEEE Transactions on Reliability, v. 67, n. 3, p. 1159-1175, set. 2018. DOI 10.1109/TR.2018.2839339.",
        "OWASP.": "OWASP. OWASP Top 10 – 2017: the ten most critical web application security risks. [S. l.]: OWASP Foundation, 2017. Disponível em: https://owasp.org/www-project-top-ten/2017/. Acesso em: 15 out. 2026.",
        "PACKER,": "PACKER, A. L. SciELO, Biblioteca Eletrônica Científica Online. In: ICSU PRESS WORKSHOP, 1998, Oxford. Oxford: Keble College, 1998. Disponível em: https://wp.scielo.org/wp-content/uploads/PACKER-A.L.-SciELO-Biblioteca-Eletr%C3%B4nica-Cient%C3%ADfica-Online.pdf. Acesso em: 15 out. 2026.",
        "PALMA,": "PALMA, F.; GONZALEZ-HUERTA, J.; FOUNI, M.; MOHA, N.; TREMBLAY, G.; GUÉHÉNEUC, Y.-G. Semantic analysis of RESTful APIs for the detection of linguistic patterns and antipatterns. International Journal of Cooperative Information Systems, v. 26, n. 2, 1742001, 2017. DOI 10.1142/S0218843017420011.",
        "PRATES,": "PRATES, R. O.; BARBOSA, S. D. J. Avaliação de interfaces de usuário: conceitos e métodos. In: JORNADAS DE ATUALIZAÇÃO EM INFORMÁTICA, 22., 2003, Campinas. Livro texto [...]. Porto Alegre: SBC, 2003. v. 2, p. 245-293.",
        "RAMÍREZ,": "RAMÍREZ, S. FastAPI: documentação oficial. [S. l.], [s. d.]. Software e documentação. Disponível em: https://fastapi.tiangolo.com/. Acesso em: 15 out. 2026.",
        "TERUEL,": "TERUEL, E. C. HTML5: guia prático. 2. ed. rev., atual. e ampl. São Paulo: Érica, 2014.",
        "VAN ROSSUM,": "VAN ROSSUM, G.; DE BOER, J. Interactively testing remote servers using the Python programming language. CWI Quarterly, Amsterdam, v. 4, n. 4, p. 283-303, dez. 1991.",
        "W3C. HTML5": "W3C. HTML5: a vocabulary and associated APIs for HTML and XHTML. W3C Recommendation, 28 out. 2014. Disponível em: https://www.w3.org/TR/html5/. Acesso em: 15 out. 2026.",
        "W3C. [": "W3C. Media Queries. W3C Recommendation, 19 jun. 2012. Disponível em: https://www.w3.org/TR/2012/REC-css3-mediaqueries-20120619/. Acesso em: 15 out. 2026.",
        "W3C. Web Content": "W3C. Web Content Accessibility Guidelines (WCAG) 2.1. W3C Recommendation, 5 jun. 2018. Disponível em: https://www.w3.org/TR/WCAG21/. Acesso em: 15 out. 2026.",
    }
    for prefix, value in references.items():
        replace_paragraph(document, prefix, value)

    remove_paragraph(find_paragraph(document, "PATHAK,"))

    appendix = find_paragraph(document, "APÊNDICE — NOTAS PARA O AUTOR")
    paragraphs = document.paragraphs
    start = next(i for i, paragraph in enumerate(paragraphs) if paragraph._p is appendix._p)
    for paragraph in list(paragraphs[start:]):
        remove_paragraph(paragraph)

    tmp = DOCX.with_suffix(".tmp.docx")
    document.save(tmp)
    os.replace(tmp, DOCX)


if __name__ == "__main__":
    main()
