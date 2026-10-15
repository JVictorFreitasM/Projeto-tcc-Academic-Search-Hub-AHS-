# ruff: noqa: E501, RUF001
"""Inclui diagramas e resultados medidos no DOCX oficial do TCC."""

from __future__ import annotations

import os
import shutil
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parents[1]
DOCX = ROOT / "tcc" / "TCC_COMPLETO_Joao_Victor.docx"
BACKUP = ROOT / "tcc" / "TCC_COMPLETO_Joao_Victor_ANTES_RESULTADOS_20261015.docx"
DIAGRAM_DIR = ROOT / "tcc" / "diagramas"


def find_paragraph(document: Document, prefix: str) -> Paragraph:
    matches = [p for p in document.paragraphs if p.text.startswith(prefix)]
    if len(matches) != 1:
        raise RuntimeError(
            f"Esperava um parágrafo iniciado por {prefix!r}; encontrei {len(matches)}"
        )
    return matches[0]


def find_first_paragraph(document: Document, prefix: str) -> Paragraph:
    for paragraph in document.paragraphs:
        if paragraph.text.startswith(prefix):
            return paragraph
    raise RuntimeError(f"Nenhum parágrafo iniciado por {prefix!r}")


def clone_after(template: Paragraph, anchor: Paragraph, text: str) -> Paragraph:
    element = deepcopy(template._p)
    anchor._p.addnext(element)
    paragraph = Paragraph(element, anchor._parent)
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in list(paragraph.runs[1:]):
            run._element.getparent().remove(run._element)
    else:
        paragraph.add_run(text)
    return paragraph


def insert_body_after(template: Paragraph, anchor: Paragraph, text: str) -> Paragraph:
    return clone_after(template, anchor, text)


def insert_figure(
    document: Document,
    anchor: Paragraph,
    caption_template: Paragraph,
    source_template: Paragraph,
    number: int,
    title: str,
    filename: str,
    *,
    page_break: bool = True,
) -> Paragraph:
    caption = clone_after(
        caption_template,
        anchor,
        f"Figura {number} – {title}",
    )
    if page_break:
        caption.paragraph_format.page_break_before = True
    caption.paragraph_format.keep_with_next = True

    picture = document.add_paragraph()
    picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture.paragraph_format.keep_with_next = True
    picture.add_run().add_picture(str(DIAGRAM_DIR / filename), width=Cm(15.5))
    caption._p.addnext(picture._p)

    source = clone_after(
        source_template,
        picture,
        "Fonte: Elaborado pelo Autor (2026).",
    )
    return source


def insert_table_after(
    document: Document,
    anchor: Paragraph,
    headers: tuple[str, ...],
    rows: tuple[tuple[str, ...], ...],
):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for cell, value in zip(table.rows[0].cells, headers, strict=True):
        cell.text = value
    for values in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, values, strict=True):
            cell.text = value
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(8)
                    if row_index == 0:
                        run.bold = True
    anchor._p.addnext(table._tbl)
    return table


def replace_text(document: Document, prefix: str, old: str, new: str) -> None:
    paragraph = find_paragraph(document, prefix)
    if old not in paragraph.text:
        raise RuntimeError(f"Texto {old!r} não encontrado em {prefix!r}")
    paragraph.text = paragraph.text.replace(old, new)


def main() -> None:
    missing = [p for p in DIAGRAM_DIR.glob("*.png") if not p.is_file()]
    if missing:
        raise RuntimeError(f"Diagramas ausentes: {missing}")
    if not BACKUP.exists():
        shutil.copy2(DOCX, BACKUP)

    document = Document(DOCX)
    body_template = find_paragraph(document, "O escopo do trabalho")
    heading_template = find_paragraph(document, "4.6 SEGURANÇA")
    caption_template = find_paragraph(document, "Figura 1")
    source_template = find_first_paragraph(document, "Fonte: Elaborado pelo Autor")

    # Renumera as capturas já presentes, reservando 1 a 5 para os diagramas.
    figure_updates = (
        ("Figura 6 – Seleção", "Figura 6", "Figura 11"),
        ("Figura 5 – Favoritos", "Figura 5", "Figura 10"),
        ("Figura 4 – Resultados em largura", "Figura 4", "Figura 9"),
        ("Figura 3 – Painel", "Figura 3", "Figura 8"),
        ("Figura 2 – Resultados", "Figura 2", "Figura 7"),
        ("Figura 1 – Resultados", "Figura 1", "Figura 6"),
        ("A aplicação web reuniu", "Figura 1", "Figura 6"),
        ("A aplicação web reuniu", "Figura 2", "Figura 7"),
        ("Os resultados puderam", "Figura 3", "Figura 8"),
        ("O layout foi construído", "Figura 4", "Figura 9"),
        ("O usuário cadastrado", "Figura 5", "Figura 10"),
        ("A exportação gerou", "Figura 6", "Figura 11"),
    )
    for prefix, old, new in figure_updates:
        replace_text(document, prefix, old, new)

    method_anchor = find_paragraph(document, "Para SciELO, BDTD")
    method_anchor = insert_body_after(
        body_template,
        method_anchor,
        "A modelagem do artefato foi documentada com diagramas de casos de uso, sequência, entidade-relacionamento e classes. O diagrama de casos de uso delimita os atores e as funções disponíveis; os dois diagramas de sequência descrevem a busca unificada e o fluxo de autenticação e favoritos; o modelo entidade-relacionamento representa a persistência; e o diagrama de classes resume os principais componentes do servidor. As Figuras 1 a 5 apresentam esses modelos.",
    )
    diagrams = (
        (1, "Diagrama de casos de uso", "01-casos-de-uso.png"),
        (2, "Diagrama de sequência da busca unificada", "02-sequencia-busca.png"),
        (
            3,
            "Diagrama de sequência da autenticação e do registro de favorito",
            "03-sequencia-autenticacao-favorito.png",
        ),
        (4, "Diagrama entidade-relacionamento", "04-entidade-relacionamento.png"),
        (5, "Diagrama de classes", "05-classes.png"),
    )
    for number, title, filename in diagrams:
        method_anchor = insert_figure(
            document,
            method_anchor,
            caption_template,
            source_template,
            number,
            title,
            filename,
        )

    method_anchor = insert_body_after(
        body_template,
        method_anchor,
        "Para avaliar o comportamento da busca, foi definido um cenário controlado com perfil fictício de estudante e cinco consultas: “letramento digital”, “acessibilidade web”, “segurança em aplicações web”, “repositórios institucionais” e “educação a distância”. Não houve participante real. Foram realizadas três execuções por consulta para medir tempo de resposta, estabilidade das fontes, resultados exibidos e duplicatas fundidas. Em uma execução adicional por consulta, analisou-se a primeira página quanto à aderência lexical e à presença de autoria, ano, resumo, DOI, ligação para PDF e idioma. A aderência lexical significou somente a presença de pelo menos uma palavra relevante da consulta no título ou no resumo; portanto, não substituiu julgamento humano de relevância.",
    )

    opening = find_paragraph(document, "Este capítulo apresenta")
    opening.text = (
        "Este capítulo apresenta a aplicação desenvolvida e os resultados observados na sua verificação. "
        "Os dados foram produzidos pelo próprio projeto entre 8 e 15 de outubro de 2026, período comprovado pelos registros disponíveis, e os totais informados pelas fontes refletem o momento de cada consulta. "
        "Além dos testes técnicos e da inspeção de usabilidade, foi executado um cenário controlado sem participante real para examinar tempo de resposta, etapas do procedimento, completude dos metadados, aderência lexical, fusão de duplicatas e estabilidade das fontes."
    )

    replace_text(document, "O servidor foi escrito", "Ramírez (2018)", "Ramírez (s. d.)")

    results_anchor = find_paragraph(document, "O conjunto da aplicação")
    results_anchor = clone_after(
        heading_template,
        results_anchor,
        "4.7 ANÁLISE DO CENÁRIO CONTROLADO DE BUSCA",
    )
    results_anchor = insert_body_after(
        body_template,
        results_anchor,
        "No cenário controlado, as cinco consultas foram submetidas à aplicação três vezes, totalizando 15 execuções. Como mostra o Quadro 3, os tempos médios ficaram entre 899 e 2.379 milissegundos. A aplicação fundiu duplicatas em quatro dos cinco temas: foram 9 ocorrências fundidas nas três execuções de “letramento digital”, 9 em “acessibilidade web”, 2 em “segurança em aplicações web” e 3 em “repositórios institucionais”. Em “educação a distância”, nenhuma duplicata foi detectada nas páginas analisadas. Esses valores descrevem o ambiente e o momento da medição, não um desempenho garantido.",
    )
    table_caption = clone_after(
        caption_template,
        results_anchor,
        "Quadro 3 – Desempenho e resultados do cenário controlado",
    )
    table = insert_table_after(
        document,
        table_caption,
        (
            "Consulta",
            "Tempo médio",
            "Exibidos (média)",
            "Duplicatas fundidas",
        ),
        (
            ("Letramento digital", "1.404 ms", "17,0", "9"),
            ("Acessibilidade web", "1.508 ms", "17,0", "9"),
            ("Segurança em aplicações web", "2.379 ms", "17,7", "2"),
            ("Repositórios institucionais", "899 ms", "12,0", "3"),
            ("Educação a distância", "1.739 ms", "15,0", "0"),
        ),
    )
    table_source = clone_after(
        source_template,
        table_caption,
        "Fonte: Elaborado pelo Autor (2026), com dados de três execuções por consulta.",
    )
    table._tbl.addnext(table_source._p)
    results_anchor = insert_body_after(
        body_template,
        table_source,
        "A estabilidade diferiu entre as fontes, conforme o Quadro 4. BDTD, OpenAlex e o índice local da UFPE responderam corretamente nas 15 execuções. O CORE entregou dados em 8 execuções e limitou a taxa de requisições em 7. SciELO, UESPI, UFPI e os quatro repositórios mantidos apenas como candidatos estavam intencionalmente desabilitados e, por isso, não foram tratados como falhas operacionais. A medição evidencia uma propriedade positiva da arquitetura: a falha parcial do CORE não impediu que a busca devolvesse resultados das demais fontes. Em contrapartida, a variação externa reduz a previsibilidade do conjunto recuperado.",
    )
    table_caption = clone_after(
        caption_template,
        results_anchor,
        "Quadro 4 – Estabilidade das fontes operacionais em 15 execuções",
    )
    table = insert_table_after(
        document,
        table_caption,
        ("Fonte", "Respostas com dados", "Ocorrências sem dados", "Duração média"),
        (
            ("BDTD", "15", "0", "333 ms"),
            ("OpenAlex", "15", "0", "928 ms"),
            ("UFPE (índice local)", "15", "0", "22 ms"),
            ("CORE", "8", "7 (limite de requisições)", "1.141 ms*"),
        ),
    )
    table_source = clone_after(
        source_template,
        table_caption,
        "Fonte: Elaborado pelo Autor (2026). *Média das respostas que informaram duração.",
    )
    table._tbl.addnext(table_source._p)
    results_anchor = insert_body_after(
        body_template,
        table_source,
        "A análise de completude da primeira página mostrou autoria entre 94,1% e 100%, ano em 100%, resumo entre 89,5% e 100% e idioma entre 88,2% e 100%. A disponibilidade de identificadores e acesso direto foi menor: DOI entre 21,1% e 26,7% e ligação para PDF entre 21,1% e 41,2%. A aderência lexical foi de 100% nas cinco consultas, dentro do critério objetivo adotado. Assim, os resultados foram suficientemente descritos para triagem por título, autoria, ano e resumo, mas frequentemente exigiram abrir a página de origem para localizar o texto completo ou confirmar o identificador.",
    )
    results_anchor = insert_body_after(
        body_template,
        results_anchor,
        "Na simulação do procedimento convencional, o mesmo texto de cada consulta teria de ser submetido separadamente às quatro fontes operacionais avaliadas — CORE, OpenAlex, BDTD e índice UFPE —, seguido de cópia e consolidação dos resultados. Isso corresponde a 20 submissões para cinco temas. Na aplicação foram necessárias 5 submissões, com consolidação automática, uma redução de 75% nessa etapa repetitiva. O achado resolve o aspecto operacional central do problema de pesquisa para as fontes efetivamente integradas: uma consulta substitui quatro submissões separadas e a aplicação trata a unificação e parte das duplicatas. Contudo, essa contagem não mede tempo humano, satisfação, facilidade percebida ou qualidade semântica; tais dimensões exigem estudo posterior com participantes.",
    )
    insert_body_after(
        body_template,
        results_anchor,
        "A colheita integral usada no índice local da UFPE processou 12.381 registros: 12.366 permaneceram ativos e 15 foram marcados como removidos. Entre os ativos, 12.365 continham resumo, todos continham ano e apenas 7 apresentavam DOI. O resultado confirma a viabilidade técnica da coleta OAI-PMH e, ao mesmo tempo, mostra que a qualidade dos campos depende do repositório de origem."
    )

    conclusion = find_paragraph(document, "O objetivo geral deste trabalho")
    conclusion.text = (
        "O objetivo geral deste trabalho foi desenvolver uma aplicação web que centralizasse a busca por trabalhos acadêmicos em múltiplas fontes de acesso aberto. O núcleo do problema foi resolvido para CORE, OpenAlex, BDTD e para o repositório da UFPE colhido no índice local: o usuário submeteu cada tema uma vez, recebeu resultados normalizados em uma lista única e contou com filtros, ordenação, fusão de duplicatas, favoritos e exportação. O cenário controlado confirmou o funcionamento integrado, a tolerância à indisponibilidade parcial e a redução de 20 para 5 submissões na comparação procedimental. O objetivo original foi cumprido em parte porque SciELO, UESPI e UFPI não puderam ser ativados nas condições encontradas."
    )
    limitations = find_paragraph(document, "Também não foi medido")
    limitations.text = (
        "Foi medido o tempo de resposta da aplicação, mas não o tempo de pessoas pesquisando manualmente; portanto, o trabalho não afirma ganho de tempo humano. A comparação procedimental foi uma simulação sem participante real. A avaliação de usabilidade foi feita por inspeção, sem usuários, e a de acessibilidade não incluiu pessoas com deficiência nem tecnologias assistivas. A aderência lexical não equivale a relevância julgada por leitor, a primeira página não representa todo o universo recuperado e os resultados das fontes variam ao longo do tempo. Essas limitações preservam a validade dos achados técnicos sem generalizá-los para a experiência humana."
    )

    tmp = DOCX.with_suffix(".finalizando.docx")
    document.save(tmp)
    os.replace(tmp, DOCX)


if __name__ == "__main__":
    main()
